"""Escribe un diccionario {nombre_hoja: DataFrame} a un único archivo .xlsx,
con formato mínimo (encabezado en negrita, fila de encabezado congelada,
ancho de columna aproximado)."""
from __future__ import annotations

import os

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

_LONGITUD_MAXIMA_NOMBRE_HOJA = 31  # límite duro de Excel


def _nombre_hoja_valido(nombre: str, usados: set) -> str:
    limpio = "".join(c for c in nombre if c not in r'[]:*?/\\')[:_LONGITUD_MAXIMA_NOMBRE_HOJA]
    candidato = limpio or "Hoja"
    sufijo = 1
    while candidato in usados:
        sufijo += 1
        recorte = _LONGITUD_MAXIMA_NOMBRE_HOJA - len(str(sufijo)) - 1
        candidato = f"{limpio[:recorte]}_{sufijo}"
    usados.add(candidato)
    return candidato


def exportar(hojas: dict[str, pd.DataFrame], ruta_salida: str) -> str:
    carpeta = os.path.dirname(ruta_salida)
    if carpeta:
        os.makedirs(carpeta, exist_ok=True)

    usados: set = set()
    with pd.ExcelWriter(ruta_salida, engine="openpyxl") as writer:
        for nombre, df in hojas.items():
            if df is None:
                continue
            nombre_hoja = _nombre_hoja_valido(nombre, usados)
            incluir_indice = df.index.name is not None
            df.to_excel(writer, sheet_name=nombre_hoja, index=incluir_indice)
            _formatear_hoja(writer.sheets[nombre_hoja], df, incluir_indice)

    return ruta_salida


def _formatear_hoja(hoja, df: pd.DataFrame, incluir_indice: bool) -> None:
    fila_encabezado = 1
    columna_inicial = 2 if incluir_indice else 1
    for columna in range(columna_inicial, hoja.max_column + 1):
        celda = hoja.cell(row=fila_encabezado, column=columna)
        celda.font = Font(bold=True)

    hoja.freeze_panes = hoja.cell(row=2, column=columna_inicial)

    columnas = ([df.index.name or ""] if incluir_indice else []) + [str(c) for c in df.columns]
    for i, nombre_columna in enumerate(columnas, start=1):
        letra = get_column_letter(i)
        try:
            valores = df.index if (incluir_indice and i == 1) else df.iloc[:, i - (2 if incluir_indice else 1)]
            ancho_datos = max((len(str(v)) for v in valores), default=0)
        except Exception:
            ancho_datos = 0
        hoja.column_dimensions[letra].width = min(60, max(10, len(nombre_columna) + 2, ancho_datos + 2))
